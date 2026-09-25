"""Render a :class:`ReportDocument` as PDF (HTML → WeasyPrint) or DOCX (python-docx)."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor
from jinja2 import Environment, FileSystemLoader, select_autoescape

from app.services.reports.content import (
    BulletList,
    KeyValues,
    Paragraph,
    ReportDocument,
    Subheading,
    Table,
)

TEMPLATES = Path(__file__).parent / "templates"
MUTED = RGBColor(0x5C, 0x5C, 0x5C)


class ReportRenderingError(Exception):
    """The report could not be rendered (the message is shown to planners)."""


_environment = Environment(
    loader=FileSystemLoader(TEMPLATES),
    autoescape=select_autoescape(["html"]),
    trim_blocks=True,
    lstrip_blocks=True,
)
_environment.filters["kind"] = lambda block: type(block).__name__


def render_html(doc: ReportDocument) -> str:
    """The report as a standalone, print-styled HTML page (the source of the PDF)."""
    return _environment.get_template("report.html").render(doc=doc)


def render_pdf(doc: ReportDocument) -> bytes:
    """Render the report to PDF with WeasyPrint.

    WeasyPrint needs the Pango system libraries (installed in the Docker image). Without them
    its import fails with ``OSError``, reported as a :class:`ReportRenderingError`.
    """
    try:
        from weasyprint import HTML
    except (OSError, ImportError) as exc:
        raise ReportRenderingError(
            "PDF generation is not available on this server (WeasyPrint's system libraries "
            "are missing). Generate a DOCX report instead, or ask an administrator."
        ) from exc
    return HTML(string=render_html(doc), base_url=str(TEMPLATES)).write_pdf()


# -------------------------------------------------------------------------------- DOCX


def _docx_table(document: DocxDocument, block: Table) -> None:
    table = document.add_table(rows=1, cols=len(block.columns))
    table.style = "Light List"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, column in enumerate(block.columns):
        cell = table.rows[0].cells[index]
        cell.text = column
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        if index in block.numeric:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in block.rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
            if index in block.numeric:
                cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if block.widths:
        usable = Mm(174)  # A4 width minus margins
        total = sum(block.widths)
        for index, width in enumerate(block.widths):
            for cell in table.columns[index].cells:
                cell.width = int(usable * width / total)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    document.add_paragraph()


def render_docx(doc: ReportDocument) -> bytes:
    """Render the report as a Word document with the same structure as the PDF."""
    document = DocxDocument()
    for section in document.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(18)
        section.top_margin, section.bottom_margin = Mm(20), Mm(18)
        section.header.paragraphs[0].text = f"NLP-RS · {doc.subtitle}"
        section.header.paragraphs[0].runs[0].font.size = Pt(8)
        section.header.paragraphs[0].runs[0].font.color.rgb = MUTED
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    document.add_heading(doc.title, level=0)
    subtitle = document.add_paragraph(doc.subtitle)
    subtitle.runs[0].font.size = Pt(13)
    meta = document.add_table(rows=0, cols=2)
    for label, value in doc.meta:
        cells = meta.add_row().cells
        cells[0].text, cells[1].text = label, value
        cells[0].paragraphs[0].runs[0].font.color.rgb = MUTED
    document.add_paragraph()

    for number, section in enumerate(doc.sections, start=1):
        document.add_heading(f"{number}. {section.title}", level=1)
        for block in section.blocks:
            if isinstance(block, Paragraph):
                paragraph = document.add_paragraph(block.text)
                if block.muted:
                    paragraph.runs[0].font.color.rgb = MUTED
            elif isinstance(block, Subheading):
                document.add_heading(block.text, level=2)
            elif isinstance(block, KeyValues):
                table = document.add_table(rows=0, cols=2)
                for label, value in block.items:
                    cells = table.add_row().cells
                    cells[0].text, cells[1].text = label, value
                document.add_paragraph()
            elif isinstance(block, BulletList):
                for item in block.items:
                    document.add_paragraph(item, style="List Bullet")
            elif isinstance(block, Table):
                _docx_table(document, block)

    disclaimer = document.add_paragraph(doc.disclaimer)
    disclaimer.runs[0].italic = True
    disclaimer.runs[0].font.color.rgb = MUTED

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
