"""Text extraction from PDF, DOCX and TXT files.

Every parser returns a :class:`ParsedDocument`: a list of pages (PDF) or a single page
(DOCX/TXT, which have no reliable page numbers).
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from app.models.enums import FileType

# Average extractable characters per page below which a PDF is treated as scanned images.
MIN_CHARS_PER_PAGE = 50
SCANNED_PDF_MESSAGE = "No extractable text (scanned PDF). OCR is not supported."


class ParseError(ValueError):
    """The file could not be turned into text; ``str(exc)`` is shown to the user."""


@dataclass(frozen=True)
class ParsedDocument:
    """Raw extracted text.

    ``pages`` holds one string per PDF page, or a single string for DOCX/TXT.
    ``page_count`` is None for formats without pages.
    """

    pages: list[str]
    page_count: int | None


def parse_pdf(content: bytes) -> ParsedDocument:
    """Extract the text of each page with PyMuPDF; reject encrypted and scanned PDFs."""
    import pymupdf

    try:
        document = pymupdf.open(stream=content, filetype="pdf")
    except Exception as exc:  # PyMuPDF raises several exception types for corrupt files
        raise ParseError("The PDF could not be opened; it may be damaged.") from exc
    with document:
        if document.needs_pass:
            raise ParseError("The PDF is password-protected.")
        pages = [page.get_text("text") for page in document]
    if not pages:
        raise ParseError("The PDF has no pages.")
    average = sum(len(page.strip()) for page in pages) / len(pages)
    if average < MIN_CHARS_PER_PAGE:
        raise ParseError(SCANNED_PDF_MESSAGE)
    return ParsedDocument(pages=pages, page_count=len(pages))


def parse_docx(content: bytes) -> ParsedDocument:
    """Extract paragraphs and table rows in document order with python-docx.

    Each table row becomes one line with cells separated by " | ", ending with a full stop so
    that sentence splitting does not merge rows.
    """
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise ParseError("The DOCX file could not be opened; it may be damaged.") from exc

    lines: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            lines.append(block.text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = _unique_cells(cell.text.strip() for cell in row.cells)
                if cells:
                    row_text = " | ".join(cells)
                    lines.append(row_text if row_text.endswith((".", "!", "?")) else row_text + ".")
            lines.append("")
    return ParsedDocument(pages=["\n".join(lines)], page_count=None)


def _unique_cells(cells: object) -> list[str]:
    """Non-empty cell texts, dropping repeats caused by merged cells."""
    result: list[str] = []
    for text in cells:  # type: ignore[attr-defined]
        if text and (not result or result[-1] != text):
            result.append(text)
    return result


def decode_text(content: bytes) -> str:
    """Decode a text file.

    Tries UTF-8 (with or without BOM), then Windows-1252 (the usual encoding of non-UTF-8
    English documents saved on Windows), then the encoding detected by charset-normalizer,
    and finally Latin-1, which accepts any bytes.
    """
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    from charset_normalizer import from_bytes

    best = from_bytes(content).best()
    if best is not None and best.encoding:
        return str(best)
    return content.decode("latin-1")


def parse_txt(content: bytes) -> ParsedDocument:
    return ParsedDocument(pages=[decode_text(content)], page_count=None)


def parse_file(file_type: FileType, content: bytes) -> ParsedDocument:
    """Dispatch to the parser for ``file_type`` and reject documents without text."""
    parser = {FileType.PDF: parse_pdf, FileType.DOCX: parse_docx, FileType.TXT: parse_txt}[
        FileType(file_type)
    ]
    parsed = parser(content)
    if not any(page.strip() for page in parsed.pages):
        raise ParseError("The document contains no text.")
    return parsed
