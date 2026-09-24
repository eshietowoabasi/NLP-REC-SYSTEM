"""Builders for SYNTHETIC test documents.

The text below is invented for tests (fictional employers and adverts). It is not real corpus
data and must never be presented as such.
"""

from __future__ import annotations

import io

import docx
import pymupdf

JOB_AD_TEXT = (
    "Synthetic Test Employer Ltd is hiring a backend engineer. The engineer will design RESTful "
    "APIs with Python and Django. Experience with AWS and Docker is required. You will build "
    "data pipelines using Apache Spark. Knowledge of Kubernetes and cloud security is an "
    "advantage."
)
SECURITY_TEXT = (
    "The security operations centre monitors threats with Splunk. Analysts perform penetration "
    "testing and incident response. CompTIA Security+ or CISSP certification is preferred."
)
FINTECH_TEXT = (
    "Our fintech team integrates Paystack and Flutterwave payment gateways. Engineers write "
    "TypeScript services on Node.js and deploy them to Google Cloud. Monitoring uses Grafana."
)


def make_pdf(pages: list[str], header: str | None = None, footer: bool = False) -> bytes:
    """A PDF with one text box per page, optional repeated header and 'Page n of N' footer."""
    document = pymupdf.open()
    for number, text in enumerate(pages, start=1):
        page = document.new_page()
        if header:
            page.insert_text((72, 50), header)
        page.insert_textbox(pymupdf.Rect(72, 90, 520, 760), text, fontsize=11)
        if footer:
            page.insert_text((72, 810), f"Page {number} of {len(pages)}")
    content = document.tobytes()
    document.close()
    return content


def make_blank_pdf(pages: int = 2) -> bytes:
    """A PDF whose pages have no text layer, like a scanned document."""
    document = pymupdf.open()
    for _ in range(pages):
        page = document.new_page()
        page.draw_rect(pymupdf.Rect(72, 72, 300, 300), color=(0, 0, 0), fill=(0.8, 0.8, 0.8))
    content = document.tobytes()
    document.close()
    return content


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    """A DOCX with the given paragraphs, then an optional table."""
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
